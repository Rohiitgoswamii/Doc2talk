
from io import BytesIO

import pymupdf
from docx import Document
from fastapi import HTTPException


def extract_pdf_text(file_bytes: bytes) -> str:
    try:
        with pymupdf.open(
            stream=file_bytes,
            filetype="pdf",
        ) as pdf:

            if pdf.needs_pass:
                raise HTTPException(
                    status_code=422,
                    detail="Password-protected PDFs are not supported.",
                )

            text_parts = [
                page.get_text()
                for page in pdf
            ]

            return "\n".join(text_parts).strip()

    except HTTPException:
        raise

    except Exception as exc:
        raise HTTPException(
            status_code=422,
            detail="Could not read the PDF. It may be invalid or corrupted.",
        ) from exc


def extract_docx_text(file_bytes: bytes) -> str:
    try:
        document = Document(BytesIO(file_bytes))

        text_parts = [
            paragraph.text
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ]

        for table in document.tables:
            for row in table.rows:
                cells = [
                    cell.text.strip()
                    for cell in row.cells
                ]
                text_parts.append(" | ".join(cells))

        return "\n".join(text_parts).strip()

    except Exception as exc:
        raise HTTPException(
            status_code=422,
            detail="Could not read the DOCX. It may be invalid or corrupted.",
        ) from exc


def extract_document_text(
    file_bytes: bytes,
    extension: str,
) -> str:

    if extension == ".pdf":
        return extract_pdf_text(file_bytes)

    if extension == ".docx":
        return extract_docx_text(file_bytes)

    raise HTTPException(
        status_code=415,
        detail="Only PDF and DOCX files are supported.",
    )


def chunk_text(
    text: str,
    chunk_size: int = 1000,
    overlap: int = 150,
) -> list[str]:

    if chunk_size <= 0:
        raise ValueError("chunk_size must be positive")

    if overlap < 0 or overlap >= chunk_size:
        raise ValueError(
            "overlap must be non-negative and smaller than chunk_size"
        )

    chunks = []
    start = 0

    while start < len(text):
        end = start + chunk_size
        chunk = text[start:end].strip()

        if chunk:
            chunks.append(chunk)

        start += chunk_size - overlap

    return chunks
